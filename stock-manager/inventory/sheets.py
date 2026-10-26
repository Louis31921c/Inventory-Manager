import csv
import io
from datetime import date, datetime

from . import db
from .models import normalize_name, parse_date

INVENTORY_FIELDS = ("article", "quantity", "supplier", "delivery_date", "order_date", "site",
                    "work_item", "designation", "backorder")
CATALOG_FIELDS = ("article", "supplier", "reference", "designation")

HINTS = {
    "article": ("article", "libelle", "designation article", "item", "ref article", "produit"),
    "quantity": ("quantity", "qty", "quantite", "qte", "nombre", "nb"),
    "supplier": ("supplier", "fournisseur", "vendor", "marque"),
    "delivery_date": ("delivery", "livraison", "date livraison", "received", "date"),
    "order_date": ("order date", "commande", "date commande", "ordered"),
    "site": ("site", "chantier", "project", "projet"),
    "work_item": ("work item", "ouvrage", "task", "lot"),
    "designation": ("designation", "description", "detail", "libelle long"),
    "reference": ("reference", "ref", "code", "sku", "ean"),
    "backorder": ("backorder", "reliquat", "missing"),
}

PREVIEW_ROWS = 8
MAX_ROWS = 20000


def _rows_from_csv(raw):
    text = raw.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), dialect)
    return [row for row in reader if any(str(cell).strip() for cell in row)]


def _rows_from_xlsx(raw):
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = book.active
    rows = []
    for line in sheet.iter_rows(values_only=True):
        if line is None or not any(cell is not None and str(cell).strip() for cell in line):
            continue
        rows.append(["" if cell is None else cell for cell in line])
        if len(rows) > MAX_ROWS:
            break
    book.close()
    return rows


def read_table(raw, media_type):
    rows = _rows_from_xlsx(raw) if "sheet" in (media_type or "") else _rows_from_csv(raw)
    if not rows:
        raise ValueError("That file has no rows.")
    header = [str(cell).strip() for cell in rows[0]]
    body = [[cell for cell in row] for row in rows[1:]]
    return header, body


def guess_mapping(header, target="inventory"):
    fields = INVENTORY_FIELDS if target == "inventory" else CATALOG_FIELDS
    mapping = {}
    for field in fields:
        for index, name in enumerate(header):
            plain = str(name).strip().lower()
            if not plain or index in mapping.values():
                continue
            if plain == field or any(hint in plain for hint in HINTS.get(field, ())):
                mapping[field] = index
                break
    return mapping


def _number(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _day(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return parse_date(str(value).strip()) if value not in (None, "") else None


def _flag(value):
    return str(value).strip().lower() in ("1", "true", "yes", "y", "oui", "x", "reliquat")


def shape(header, body, mapping, target="inventory"):
    """Turn the picked columns into rows this app can store, and say what it threw out."""
    fields = INVENTORY_FIELDS if target == "inventory" else CATALOG_FIELDS
    mapping = {field: int(index) for field, index in mapping.items()
               if field in fields and str(index) != ""}
    if "article" not in mapping:
        raise ValueError("Point the article column out first.")

    kept, dropped = [], []
    for number, line in enumerate(body[:MAX_ROWS], 2):
        def cell(field):
            index = mapping.get(field)
            if index is None or index >= len(line):
                return None
            value = line[index]
            return None if value in (None, "") else value

        article = normalize_name(str(cell("article") or ""))
        if not article:
            dropped.append({"row": number, "why": "no article"})
            continue

        row = {"article": article,
               "supplier": normalize_name(str(cell("supplier") or "")) or None,
               "designation": (str(cell("designation")).strip()
                               if cell("designation") is not None else None)}
        if target == "inventory":
            quantity = _number(cell("quantity"))
            if quantity is None:
                dropped.append({"row": number, "why": "no quantity"})
                continue
            row.update(quantity=quantity,
                       delivery_date=_day(cell("delivery_date")),
                       order_date=_day(cell("order_date")),
                       site=normalize_name(str(cell("site") or "")) or None,
                       work_item=normalize_name(str(cell("work_item") or "")) or None,
                       backorder=_flag(cell("backorder")))
            if row["delivery_date"] is None:
                dropped.append({"row": number, "why": "no delivery date"})
                continue
        else:
            row.update(reference=(str(cell("reference")).strip()
                                  if cell("reference") is not None else None))
            if not row["supplier"]:
                dropped.append({"row": number, "why": "no supplier"})
                continue
        kept.append(row)
    return kept, dropped


def store_inventory(db_path, rows, user=None):
    """Imported lines hang off one note of their own, so they can be found again."""
    if not rows:
        return {"note_id": None, "lines": 0}
    note_id = db.write(db_path, """
            INSERT INTO notes (image_sha256, image_path, confirmed_by)
            VALUES (NULL, NULL, ?) RETURNING note_id
            """, (user,))[0]
    for line_no, row in enumerate(rows, 1):
        db.write(db_path, """
                INSERT INTO inventory_lines
                    (article, quantity, supplier, delivery_date, order_date, site,
                     work_item, backorder, designation, note_id, line_no)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (row["article"], row["quantity"], row["supplier"], row["delivery_date"],
                      row["order_date"], row["site"], row["work_item"], row["backorder"],
                      row["designation"], note_id, line_no))
    return {"note_id": note_id, "lines": len(rows)}


def store_catalog(db_path, rows):
    for row in rows:
        db.write(db_path, """
                INSERT INTO catalog (supplier, designation, reference, article, seen, updated_at)
                VALUES (?, ?, ?, ?, 1, current_timestamp)
                """, (row["supplier"], row["designation"], row["reference"], row["article"]))
    return {"lines": len(rows)}
